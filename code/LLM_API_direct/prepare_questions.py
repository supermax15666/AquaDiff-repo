from __future__ import annotations

import argparse
import json
import math
import random
import statistics
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from common import (
    SCHEMA_VERSION,
    TARGET_UNIT,
    atomic_write_json,
    atomic_write_jsonl,
    canonical_json_bytes,
    iter_jsonl,
    resolve_path,
    safe_slug,
    sha256_bytes,
    sha256_file,
    stable_hash,
    utc_now,
)


DEFAULT_INPUT_DIR = Path(__file__).resolve().parents[2] / "分谱系数据"
DEFAULT_OUTPUT = (
    Path(__file__).resolve().parent / "questions" / "four_system_all_questions.jsonl"
)
DEFAULT_GOLD_OUTPUT = (
    Path(__file__).resolve().parent / "questions" / "four_system_all_gold.evaluator.jsonl"
)
DEFAULT_ANALYSIS_SETS_OUTPUT = (
    Path(__file__).resolve().parent / "questions" / "four_system_analysis_sets.json"
)
DEFAULT_SENSITIVITY_SIZE_PER_SYSTEM = 60
FROZEN_SENSITIVITY_MEMBER_IDS_SHA256 = (
    "dae63a99271de1286be373d7956315c3b802ade1d4a840f63ace99f644c05d89"
)
FROZEN_PRIMARY_MEMBER_IDS_SHA256 = (
    "bea7c062026bed14a6c05da61f461abdbd6ce51be26e725b02f9a129c99a2579"
)
FROZEN_PRIMARY_QUESTIONS_SHA256 = (
    "605eae4eba17e9aff06c80e7e40152f1c6b4fa11fa5a8b31138faa88fd1f8414"
)
FROZEN_PRIMARY_GOLD_SHA256 = (
    "a0d0a610285e6499ed3f76c4324dfc969f1ac683083ec3e337cc7cdde2305471"
)
FROZEN_BALANCED_QUESTIONS = (
    Path(__file__).resolve().parent / "questions" / "four_system_questions.jsonl"
)
FROZEN_BALANCED_GOLD = (
    Path(__file__).resolve().parent / "questions" / "four_system_gold.evaluator.jsonl"
)
EXPECTED_WORKBOOKS = {
    "NaCl-H2O": "NaCl分谱系数据.xlsx",
    "KCl-H2O": "KCl分谱系数据.xlsx",
    "LiCl-H2O": "LiCl-H2O分谱系数据.xlsx",
    "CaCl2-H2O": "CaCl2-H2O分谱系数据.xlsx",
}
REQUIRED_COLUMNS = {
    "system",
    "temperature_K",
    "concentration_mol_L",
    "diffusion_1e9_m2_s",
    "model_group_id",
}


def finite_float(value: Any, name: str, record_id: str) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Invalid {name!r} in {record_id!r}: {value!r}") from exc
    if not math.isfinite(number):
        raise ValueError(f"Non-finite {name!r} in {record_id!r}")
    return number


def template_text(language: str, system: str, temperature: float, concentration: float) -> str:
    temperature_text = f"{temperature:.2f}"
    concentration_text = f"{concentration:.6f}"
    if language == "zh":
        return (
            f"预测 {system} 在 T = {temperature_text} K、c = {concentration_text} mol/L "
            "时的互扩散系数。返回 D = x * 10^-9 m^2/s 中的乘数 x。JSON 数值字段"
            "只能包含 x；不要返回 x * 10^-9，也不要返回以基础 SI 单位表示的 D。"
        )
    return (
        f"Predict the mutual diffusion coefficient of {system} at "
        f"T = {temperature_text} K and c = {concentration_text} mol/L. "
        "Return the multiplier x where D = x * 10^-9 m^2/s. "
        "The JSON numeric field must contain x only; do not return x * 10^-9 "
        "and do not return D in base SI units."
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Build a blinded four-system API benchmark from lineage-audited workbooks. "
            "Same-condition literature records are collapsed to a median consensus gold."
        )
    )
    parser.add_argument("--input-dir", type=Path, default=DEFAULT_INPUT_DIR)
    parser.add_argument("--sheet", default="Confirmed_Records")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--gold-output", type=Path, default=DEFAULT_GOLD_OUTPUT)
    parser.add_argument(
        "--per-system",
        type=int,
        default=0,
        help="Unique conditions per system; use 0 for every unique condition.",
    )
    parser.add_argument("--selection-seed", type=int, default=20260812)
    parser.add_argument(
        "--analysis-sets-output",
        type=Path,
        default=DEFAULT_ANALYSIS_SETS_OUTPUT,
        help="Frozen membership manifest for the all-condition primary and balanced sensitivity sets.",
    )
    parser.add_argument(
        "--sensitivity-per-system",
        type=int,
        default=DEFAULT_SENSITIVITY_SIZE_PER_SYSTEM,
        help="Balanced sensitivity-subset size per system; no extra API calls are made.",
    )
    parser.add_argument(
        "--include-multi-lineage-conditions",
        action="store_true",
        help=(
            "Allow conditions reported by several lineages. By default these are kept in "
            "the audit manifest but excluded from the balanced paper set."
        ),
    )
    parser.add_argument(
        "--include-zero-concentration",
        action="store_true",
        help="Include c=0 limiting/derived records; excluded by default from the paper set.",
    )
    parser.add_argument("--language", choices=["en", "zh"], default="en")
    parser.add_argument(
        "--evidence-level",
        default="retrospective_opened_four_system_lineage_pool",
        help="The current workbooks are not a sealed prospective holdout.",
    )
    return parser.parse_args()


def read_workbook(path: Path, sheet_name: str, expected_system: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Workbook {path.name!r} is missing sheet {sheet_name!r}")
        sheet = workbook[sheet_name]
        iterator = sheet.iter_rows(values_only=True)
        try:
            headers = [str(value).strip() if value is not None else "" for value in next(iterator)]
        except StopIteration as exc:
            raise ValueError(f"Workbook sheet is empty: {path}::{sheet_name}") from exc
        missing = sorted(REQUIRED_COLUMNS - set(headers))
        if missing:
            raise ValueError(f"Workbook {path.name!r} is missing columns: {missing}")
        records: list[dict[str, Any]] = []
        for row_number, values in enumerate(iterator, start=2):
            if not any(value is not None for value in values):
                continue
            row = dict(zip(headers, values))
            record_id = f"{path.stem}::row_{row_number:06d}"
            system = str(row.get("system") or "").strip()
            lineage = str(row.get("model_group_id") or "").strip()
            if system != expected_system:
                raise ValueError(
                    f"Unexpected system in {record_id}: {system!r}; expected {expected_system!r}"
                )
            if not lineage:
                raise ValueError(f"Missing model_group_id in {record_id}")
            temperature = finite_float(row.get("temperature_K"), "temperature_K", record_id)
            concentration = finite_float(
                row.get("concentration_mol_L"), "concentration_mol_L", record_id
            )
            diffusion = finite_float(
                row.get("diffusion_1e9_m2_s"), "diffusion_1e9_m2_s", record_id
            )
            if temperature <= 0 or concentration < 0 or diffusion <= 0:
                raise ValueError(f"Non-physical value in {record_id}")
            records.append(
                {
                    "source_record_id": record_id,
                    "source_workbook": path.name,
                    "source_row": row_number,
                    "system": system,
                    "temperature_K": temperature,
                    "concentration_mol_L": concentration,
                    "diffusion_1e9_m2_s": diffusion,
                    "origin_lineage_id": lineage,
                    "lobo_ref_no": row.get("lobo_ref_no"),
                    "original_reference": str(row.get("original_reference") or "").strip(),
                }
            )
        if not records:
            raise ValueError(f"No records found in {path}::{sheet_name}")
        return records
    finally:
        workbook.close()


def condition_key(record: dict[str, Any]) -> tuple[str, float, float]:
    return (
        str(record["system"]),
        float(record["temperature_K"]),
        float(record["concentration_mol_L"]),
    )


def aggregate_conditions(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[tuple[str, float, float], list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        grouped[condition_key(record)].append(record)
    conditions: list[dict[str, Any]] = []
    for (system, temperature, concentration), source_records in sorted(grouped.items()):
        source_records.sort(key=lambda row: row["source_record_id"])
        values = [float(row["diffusion_1e9_m2_s"]) for row in source_records]
        lineages = sorted({str(row["origin_lineage_id"]) for row in source_records})
        lineage_cluster_id = (
            lineages[0]
            if len(lineages) == 1
            else "LINEAGE_SET:" + stable_hash(lineages, 20)
        )
        condition_id = stable_hash(
            {
                "system": system,
                "temperature_K": temperature,
                "concentration_mol_L": concentration,
            },
            20,
        )
        consensus = float(statistics.median(values))
        conditions.append(
            {
                "condition_id": condition_id,
                "system": system,
                "temperature_K": temperature,
                "concentration_mol_L": concentration,
                "diffusion_1e9_m2_s": consensus,
                "reference_value_policy": "median_across_same_condition_records",
                "reference_value_count": len(values),
                "reference_value_min_1e9_m2_s": min(values),
                "reference_value_max_1e9_m2_s": max(values),
                "reference_value_relative_range": (
                    (max(values) - min(values)) / abs(consensus) if consensus else None
                ),
                "origin_lineage_id": lineage_cluster_id,
                "origin_lineage_ids": lineages,
                "source_records": source_records,
            }
        )
    return conditions


def proportional_allocations(
    available: dict[str, list[dict[str, Any]]], target: int
) -> dict[str, int]:
    total = sum(len(rows) for rows in available.values())
    if target >= total:
        return {lineage: len(rows) for lineage, rows in available.items()}
    allocations = {lineage: 0 for lineage in available}
    lineages = sorted(available)
    if target >= len(lineages):
        for lineage in lineages:
            allocations[lineage] = 1
    else:
        lineages = sorted(lineages, key=lambda key: (-len(available[key]), key))[:target]
        for lineage in lineages:
            allocations[lineage] = 1
        return allocations
    while sum(allocations.values()) < target:
        candidates = [
            lineage
            for lineage in available
            if allocations[lineage] < len(available[lineage])
        ]
        chosen = max(
            candidates,
            key=lambda lineage: (
                target * len(available[lineage]) / total - allocations[lineage],
                len(available[lineage]) - allocations[lineage],
                lineage,
            ),
        )
        allocations[chosen] += 1
    return allocations


def lineage_is_derived(lineage: str) -> bool:
    upper = lineage.upper()
    return upper.startswith("DERIVED:") or ":HARNED_1955:D0" in upper


def manifest_relative_path(target: Path, manifest_path: Path) -> str:
    try:
        return target.resolve().relative_to(manifest_path.parent.resolve()).as_posix()
    except ValueError:
        return str(target.resolve())


def jsonl_sha256(records: list[dict[str, Any]]) -> str:
    lines = [canonical_json_bytes(record) for record in records]
    payload = b"\n".join(lines) + (b"\n" if lines else b"")
    return sha256_bytes(payload)


def select_conditions(
    conditions: list[dict[str, Any]], per_system: int, seed: int
) -> list[dict[str, Any]]:
    by_system: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for condition in conditions:
        by_system[str(condition["system"])].append(condition)
    selected: list[dict[str, Any]] = []
    system_order = list(EXPECTED_WORKBOOKS)
    for system in sorted(by_system, key=lambda value: system_order.index(value)):
        rows = by_system[system]
        target = len(rows) if per_system == 0 else per_system
        if target > len(rows):
            raise ValueError(
                f"Requested {target} questions for {system}, but only {len(rows)} unique conditions exist"
            )
        by_lineage: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for row in rows:
            by_lineage[str(row["origin_lineage_id"])].append(row)
        allocations = proportional_allocations(by_lineage, target)
        for lineage, count in sorted(allocations.items()):
            candidates = list(by_lineage[lineage])
            random.Random(f"{seed}:{system}:{lineage}").shuffle(candidates)
            selected.extend(candidates[:count])
    return sorted(
        selected,
        key=lambda row: (
            system_order.index(row["system"]),
            row["origin_lineage_id"],
            row["temperature_K"],
            row["concentration_mol_L"],
            row["condition_id"],
        ),
    )


def main() -> int:
    args = parse_args()
    if args.per_system < 0:
        raise ValueError("--per-system must be zero or positive")
    if args.sensitivity_per_system <= 0:
        raise ValueError("--sensitivity-per-system must be positive")
    if args.per_system != 0:
        raise ValueError(
            "The active study plan requires --per-system 0: all eligible unique conditions "
            "form the primary API question set."
        )
    if args.include_multi_lineage_conditions or args.include_zero_concentration:
        raise ValueError(
            "The active frozen benchmark excludes multi-lineage and zero-concentration "
            "conditions. Create a separately named exploratory asset for alternative rules."
        )
    if args.language != "en":
        raise ValueError(
            "The active frozen benchmark uses the English direct-answer template."
        )
    project_dir = Path(__file__).resolve().parent
    input_dir = resolve_path(project_dir, args.input_dir)
    output_path = resolve_path(project_dir, args.output)
    gold_output_path = resolve_path(project_dir, args.gold_output)
    analysis_sets_path = resolve_path(project_dir, args.analysis_sets_output)

    all_records: list[dict[str, Any]] = []
    source_assets = []
    for expected_system, filename in EXPECTED_WORKBOOKS.items():
        path = input_dir / filename
        if not path.is_file():
            raise FileNotFoundError(f"Required four-system workbook not found: {path}")
        records = read_workbook(path, args.sheet, expected_system)
        all_records.extend(records)
        source_assets.append(
            {
                "path": str(path),
                "filename": filename,
                "sha256": sha256_file(path),
                "sheet": args.sheet,
                "record_count": len(records),
                "system": expected_system,
                "lineage_count": len({row["origin_lineage_id"] for row in records}),
            }
        )

    all_conditions = aggregate_conditions(all_records)
    eligible_conditions = [
        row
        for row in all_conditions
        if (args.include_multi_lineage_conditions or len(row["origin_lineage_ids"]) == 1)
        and (args.include_zero_concentration or float(row["concentration_mol_L"]) > 0)
        and not any(lineage_is_derived(lineage) for lineage in row["origin_lineage_ids"])
    ]
    selected = select_conditions(eligible_conditions, args.per_system, args.selection_seed)
    sensitivity_selected = select_conditions(
        eligible_conditions, args.sensitivity_per_system, args.selection_seed
    )
    template_family = f"direct_internalized_numeric_{args.language}_v1"
    questions: list[dict[str, Any]] = []
    gold_records: list[dict[str, Any]] = []
    for order, condition in enumerate(selected, start=1):
        system = str(condition["system"])
        question_id = (
            f"four-system::{safe_slug(system).lower()}::{condition['condition_id']}"
        )
        question: dict[str, Any] = {
            "schema_version": SCHEMA_VERSION,
            "question_id": question_id,
            "question_order": order,
            "base_prompt_id": question_id,
            "template_family_id": template_family,
            "knowledge_condition": "direct_closed_book",
            "language": args.language,
            "task_type": "direct_internalized_knowledge_numeric",
            "route": "model_prediction",
            "system": system,
            "temperature_K": condition["temperature_K"],
            "concentration_mol_L": condition["concentration_mol_L"],
            "target_property": "binary_aqueous_mutual_diffusion_coefficient",
            "target_unit": TARGET_UNIT,
            "prompt": template_text(
                args.language,
                system,
                float(condition["temperature_K"]),
                float(condition["concentration_mol_L"]),
            ),
            "scoring": {
                "relative_error_thresholds": [0.05, 0.10],
                "parse_or_unit_failure_counts_as_incorrect": True,
                "missing_or_refusal_counts_as_incorrect": True,
            },
            "evidence_level": args.evidence_level,
        }
        question["question_sha256"] = stable_hash(question)
        questions.append(question)

        gold: dict[str, Any] = {
            "schema_version": SCHEMA_VERSION,
            "question_id": question_id,
            "base_prompt_id": question_id,
            "condition_id": condition["condition_id"],
            "diffusion_1e9_m2_s": condition["diffusion_1e9_m2_s"],
            "target_unit": TARGET_UNIT,
            "origin_lineage_id": condition["origin_lineage_id"],
            "origin_lineage_ids": condition["origin_lineage_ids"],
            "split_role": "retrospective_condition_consensus",
            "system": system,
            "temperature_K": condition["temperature_K"],
            "concentration_mol_L": condition["concentration_mol_L"],
            "reference_value_policy": condition["reference_value_policy"],
            "reference_value_count": condition["reference_value_count"],
            "reference_value_min_1e9_m2_s": condition[
                "reference_value_min_1e9_m2_s"
            ],
            "reference_value_max_1e9_m2_s": condition[
                "reference_value_max_1e9_m2_s"
            ],
            "reference_value_relative_range": condition[
                "reference_value_relative_range"
            ],
            "source_records": condition["source_records"],
            "evidence_level": args.evidence_level,
        }
        gold["gold_record_sha256"] = stable_hash(gold)
        gold_records.append(gold)

    primary_ids = [record["question_id"] for record in questions]
    primary_id_set = set(primary_ids)
    rendered_prompts = [str(record["prompt"]) for record in questions]
    if len(rendered_prompts) != len(set(rendered_prompts)):
        raise ValueError(
            "Different source conditions collapse to the same rendered prompt precision"
        )
    primary_ids_sha256 = stable_hash(sorted(primary_ids))
    if primary_ids_sha256 != FROZEN_PRIMARY_MEMBER_IDS_SHA256:
        raise ValueError(
            "The 642-question primary membership changed unexpectedly; refusing to "
            "replace the frozen benchmark assets"
        )
    sensitivity_ids = [
        f"four-system::{safe_slug(str(condition['system'])).lower()}::{condition['condition_id']}"
        for condition in sensitivity_selected
    ]
    if len(sensitivity_ids) != len(set(sensitivity_ids)):
        raise ValueError("Balanced sensitivity subset contains duplicate question IDs")
    if not set(sensitivity_ids).issubset(primary_id_set):
        raise ValueError("Balanced sensitivity subset is not a strict subset of the primary set")
    sensitivity_ids_sha256 = stable_hash(sorted(sensitivity_ids))
    if (
        args.selection_seed == 20260812
        and args.sensitivity_per_system == DEFAULT_SENSITIVITY_SIZE_PER_SYSTEM
        and sensitivity_ids_sha256 != FROZEN_SENSITIVITY_MEMBER_IDS_SHA256
    ):
        raise ValueError(
            "The balanced 240-question membership changed unexpectedly; refusing to replace "
            "the frozen sensitivity subset"
        )
    if not FROZEN_BALANCED_QUESTIONS.is_file() or not FROZEN_BALANCED_GOLD.is_file():
        raise FileNotFoundError(
            "The pre-frozen balanced 240-question assets are required to verify sensitivity "
            "membership"
        )
    frozen_questions = {
        str(record["question_id"]): record
        for record in iter_jsonl(FROZEN_BALANCED_QUESTIONS)
    }
    frozen_gold = {
        str(record["question_id"]): record
        for record in iter_jsonl(FROZEN_BALANCED_GOLD)
    }
    if set(frozen_questions) != set(sensitivity_ids) or set(frozen_gold) != set(sensitivity_ids):
        raise ValueError("The pre-frozen 240-question files do not match the sensitivity members")
    primary_questions_by_id = {record["question_id"]: record for record in questions}
    primary_gold_by_id = {record["question_id"]: record for record in gold_records}
    for question_id in sensitivity_ids:
        frozen_question = {
            key: value
            for key, value in frozen_questions[question_id].items()
            if key not in {"question_order", "question_sha256"}
        }
        current_question = {
            key: value
            for key, value in primary_questions_by_id[question_id].items()
            if key not in {"question_order", "question_sha256"}
        }
        if frozen_question != current_question:
            raise ValueError(
                f"Frozen sensitivity question {question_id!r} changed content"
            )
        if frozen_gold[question_id] != primary_gold_by_id[question_id]:
            raise ValueError(f"Frozen sensitivity gold {question_id!r} changed content")

    generated_question_sha256 = jsonl_sha256(questions)
    generated_gold_sha256 = jsonl_sha256(gold_records)
    if generated_question_sha256 != FROZEN_PRIMARY_QUESTIONS_SHA256:
        raise ValueError("Generated primary question-file hash differs from the frozen asset")
    if generated_gold_sha256 != FROZEN_PRIMARY_GOLD_SHA256:
        raise ValueError("Generated evaluator-gold hash differs from the frozen asset")
    atomic_write_jsonl(output_path, questions)
    atomic_write_jsonl(gold_output_path, gold_records)
    manifest_path = output_path.with_name(output_path.stem + ".manifest.json")
    primary_system_counts = Counter(record["system"] for record in questions)
    sensitivity_system_counts = Counter(
        primary_questions_by_id[question_id]["system"] for question_id in sensitivity_ids
    )
    analysis_sets: dict[str, Any] = {
        "schema_version": SCHEMA_VERSION,
        "artifact_type": "llm_api_direct_analysis_sets",
        "response_reuse": True,
        "additional_api_calls_for_sensitivity_subset": 0,
        "primary_question_file": manifest_relative_path(output_path, analysis_sets_path),
        "primary_question_file_sha256": sha256_file(output_path),
        "evaluator_gold_file": manifest_relative_path(
            gold_output_path, analysis_sets_path
        ),
        "evaluator_gold_file_sha256": sha256_file(gold_output_path),
        "sets": [
            {
                "analysis_set_id": "primary_all_eligible_unique_conditions",
                "role": "primary",
                "question_count": len(primary_ids),
                "question_ids": primary_ids,
                "question_ids_sha256": primary_ids_sha256,
                "system_counts": dict(sorted(primary_system_counts.items())),
                "selection_policy": "all_eligible_unique_conditions_after_frozen_filters",
                "selection_seed": None,
            },
            {
                "analysis_set_id": "sensitivity_balanced_60_per_system",
                "role": "sensitivity",
                "question_count": len(sensitivity_ids),
                "question_ids": sensitivity_ids,
                "question_ids_sha256": sensitivity_ids_sha256,
                "system_counts": dict(sorted(sensitivity_system_counts.items())),
                "selection_policy": (
                    "lineage_coverage_then_proportional_allocation_seeded_without_replacement"
                ),
                "selection_seed": args.selection_seed,
                "per_system": args.sensitivity_per_system,
                "frozen_before_primary_api_run": True,
            },
        ],
        "frozen_sensitivity_source": {
            "question_file": manifest_relative_path(
                FROZEN_BALANCED_QUESTIONS, analysis_sets_path
            ),
            "question_file_sha256": sha256_file(FROZEN_BALANCED_QUESTIONS),
            "gold_file": manifest_relative_path(
                FROZEN_BALANCED_GOLD, analysis_sets_path
            ),
            "gold_file_sha256": sha256_file(FROZEN_BALANCED_GOLD),
            "question_ids_sha256": sensitivity_ids_sha256,
        },
    }
    analysis_sets["analysis_sets_content_sha256"] = stable_hash(analysis_sets)
    atomic_write_json(analysis_sets_path, analysis_sets)
    selected_lineage_counts: dict[str, Counter[str]] = defaultdict(Counter)
    for record in gold_records:
        selected_lineage_counts[record["system"]][record["origin_lineage_id"]] += 1
    all_system_counts = Counter(row["system"] for row in all_conditions)
    eligible_system_counts = Counter(row["system"] for row in eligible_conditions)
    selected_system_counts = Counter(row["system"] for row in questions)
    duplicate_conditions = [
        condition
        for condition in all_conditions
        if int(condition["reference_value_count"]) > 1
    ]
    manifest: dict[str, Any] = {
        "schema_version": SCHEMA_VERSION,
        "artifact_type": "four_system_question_set_manifest",
        "created_at_utc": utc_now(),
        "question_set_path": str(output_path),
        "question_set_sha256": sha256_file(output_path),
        "evaluator_gold_path": str(gold_output_path),
        "evaluator_gold_sha256": sha256_file(gold_output_path),
        "source_assets": source_assets,
        "source_record_count": len(all_records),
        "unique_condition_count": len(all_conditions),
        "same_condition_consensus_policy": "median",
        "same_condition_duplicate_group_count": len(duplicate_conditions),
        "same_condition_extra_record_count": sum(
            int(condition["reference_value_count"]) - 1
            for condition in duplicate_conditions
        ),
        "selection": {
            "per_system": args.per_system,
            "selection_seed": args.selection_seed,
            "policy": "lineage_coverage_then_proportional_allocation_seeded_without_replacement",
            "system_balance": "equal_when_per_system_is_positive",
            "include_multi_lineage_conditions": bool(args.include_multi_lineage_conditions),
            "include_zero_concentration": bool(args.include_zero_concentration),
            "zero_concentration_conditions_excluded_from_selection": (
                0
                if args.include_zero_concentration
                else sum(
                    float(condition["concentration_mol_L"]) == 0
                    for condition in all_conditions
                )
            ),
            "explicit_derived_lineage_conditions_excluded_from_selection": sum(
                any(lineage_is_derived(lineage) for lineage in condition["origin_lineage_ids"])
                for condition in all_conditions
            ),
            "multi_lineage_conditions_excluded_from_selection": (
                0
                if args.include_multi_lineage_conditions
                else sum(
                    len(condition["origin_lineage_ids"]) > 1
                    for condition in all_conditions
                )
            ),
            "eligible_unique_condition_count_after_all_filters": len(eligible_conditions),
            "confirmatory_split_role": None,
            "reason_no_confirmatory_split": (
                "The current four workbooks contain lineage grouping but no frozen, "
                "leakage-safe holdout assignments."
            ),
        },
        "language": args.language,
        "template_family_id": template_family,
        "evidence_level": args.evidence_level,
        "question_count": len(questions),
        "analysis_sets_path": str(analysis_sets_path),
        "analysis_sets_sha256": sha256_file(analysis_sets_path),
        "primary_analysis_set_id": "primary_all_eligible_unique_conditions",
        "sensitivity_analysis_set_id": "sensitivity_balanced_60_per_system",
        "sensitivity_question_count": len(sensitivity_ids),
        "sensitivity_question_ids_sha256": sensitivity_ids_sha256,
        "sensitivity_additional_api_calls": 0,
        "available_unique_condition_counts": dict(sorted(all_system_counts.items())),
        "eligible_unique_condition_counts": dict(sorted(eligible_system_counts.items())),
        "selected_system_counts": dict(sorted(selected_system_counts.items())),
        "selected_lineage_cluster_counts": {
            system: dict(sorted(counts.items()))
            for system, counts in sorted(selected_lineage_counts.items())
        },
        "gold_present_in_question_set": False,
        "gold_stored_separately": True,
        "gold_sent_to_model": False,
        "formal_confirmatory_use_allowed": False,
        "warning": (
            "The four lineage workbooks and their reference values are already available in "
            "the project. This benchmark is retrospective/exploratory, not a sealed blind "
            "or prospective external evaluation."
        ),
        "generator_path": str(Path(__file__).resolve()),
        "generator_sha256": sha256_file(Path(__file__).resolve()),
    }
    manifest["manifest_content_sha256"] = stable_hash(manifest)
    atomic_write_json(manifest_path, manifest)
    print(
        json.dumps(
            {
                "source_records": len(all_records),
                "unique_conditions": len(all_conditions),
                "questions": len(questions),
                "selected_system_counts": manifest["selected_system_counts"],
                "output": str(output_path),
                "sha256": manifest["question_set_sha256"],
                "evaluator_gold": str(gold_output_path),
                "evaluator_gold_sha256": manifest["evaluator_gold_sha256"],
                "manifest": str(manifest_path),
                "analysis_sets": str(analysis_sets_path),
                "analysis_sets_sha256": manifest["analysis_sets_sha256"],
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
