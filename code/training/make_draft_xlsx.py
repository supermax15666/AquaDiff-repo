# -*- coding: utf-8 -*-
"""把草稿 CSV 与 B2 密度查值清单转成 Excel 版（用户核对/填写用）。

- candidates_draft_v1.csv -> candidates_draft_v1.xlsx（加"核对备注"空列）
- b2_密度查值清单.csv -> b2_密度查值清单.xlsx（加"质量百分数_wt%"列，CRC 密度表按 wt% 列值）
"""
from __future__ import annotations

import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[2]
DRAFT_DIR = ROOT / "py" / "微调" / "external" / "draft"

M_KCL = 74.5513


def csv_rows(path: Path) -> list[list[str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.reader(handle))


def write_sheet(workbook: Workbook, title: str, rows: list[list[str]], widths: dict[int, int]) -> None:
    sheet = workbook.create_sheet(title)
    for row in rows:
        sheet.append(row)
    for index, cell in enumerate(sheet[1], start=1):
        cell.font = Font(bold=True)
    for column, width in widths.items():
        sheet.column_dimensions[chr(64 + column)].width = width


def main() -> None:
    draft = DRAFT_DIR / "candidates_draft_v1.csv"
    density = DRAFT_DIR / "b2_密度查值清单.csv"

    workbook = Workbook()
    workbook.remove(workbook.active)
    rows = csv_rows(draft)
    rows[0].append("核对备注（OK/修正/疑问，留空=待核）")
    for row in rows[1:]:
        row.append("")
    write_sheet(
        workbook, "草稿15行", rows,
        {1: 11, 2: 13, 3: 18, 4: 18, 5: 22, 6: 60, 7: 12, 8: 28, 9: 12, 10: 60, 11: 40},
    )
    out_draft = DRAFT_DIR / "candidates_draft_v1.xlsx"
    workbook.save(out_draft)

    workbook2 = Workbook()
    workbook2.remove(workbook2.active)
    rows2 = csv_rows(density)
    rows2[0].append("质量百分数_wt%（CRC 表按此查）")
    for row in rows2[1:]:
        try:
            m = float(row[1])
            wt = 100.0 * m * M_KCL / (1000.0 + m * M_KCL)
            row.append(f"{wt:.2f}")
        except ValueError:
            row.append("")
    write_sheet(
        workbook2, "B2密度", rows2,
        {1: 10, 2: 14, 3: 10, 4: 12, 5: 14, 6: 22, 7: 70},
    )
    out_density = DRAFT_DIR / "b2_密度查值清单.xlsx"
    workbook2.save(out_density)

    print("written:", out_draft)
    print("written:", out_density)


if __name__ == "__main__":
    main()
