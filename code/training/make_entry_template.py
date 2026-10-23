# -*- coding: utf-8 -*-
"""生成外部集录入模板 xlsx（仅表头，无示例数据行——防止示例值混入 gold）。

产物：py/微调/external/录入模板.xlsx
  - Sheet "Confirmed_Records"：仅列头（audit_freeze_external.py 默认读这个表）；
  - Sheet "填写说明"：逐列解释 + 示例（示例仅在说明表，脚本不读）。
用法：把本模板复制成 py/微调/external/candidates/<体系>候选.xlsx 再填写。
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "py" / "微调" / "external" / "录入模板.xlsx"

HEADERS = [
    "system", "temperature_K", "concentration_mol_L", "diffusion_1e9_m2_s",
    "origin_lineage_id", "original_reference", "source_type", "record_id",
    "retrieval_date", "notes",
]

EXPLANATIONS = [
    ("列名", "必填", "怎么填"),
    ("system", "✔", "NaCl-H2O / KCl-H2O / LiCl-H2O / CaCl2-H2O（四选一，一字不差）"),
    ("temperature_K", "✔", "温度，单位 K。原文 25 °C → 298.15 K"),
    ("concentration_mol_L", "✔", "浓度，单位 mol/L（mol/dm³ 等同）。原文是 mol/kg 且无法可靠换算→先不录，notes 说明"),
    ("diffusion_1e9_m2_s", "✔", "互扩散系数乘数 x，即 D = x × 10⁻⁹ m²/s。原文 10⁻⁵ cm²/s 的数值直接就是 x；原文 10⁻⁶ cm²/s → x = 数值 × 0.1"),
    ("origin_lineage_id", "✔", "新谱系名，格式 EXP:<体系缩写>:<第一作者姓>_<年份>，如 EXP:NACL:RIBEIRO_2008（缩写 NACL/KCL/LICL/CACL2）"),
    ("original_reference", "✔", "完整引文：作者、期刊、卷、页、年份、DOI"),
    ("source_type", "✔", "一律填 experimental（派生/平滑/插值/转录数据不要录）"),
    ("record_id", "可选", "该行在原文的位置，如 ribeiro2008::table1-row3（便于追溯）"),
    ("retrieval_date", "可选", "录入日期，如 2026-08-16"),
    ("notes", "可选", "表号/图号、页码、单位换算过程、任何拿不准的地方（拿不准就写这里，不要猜）"),
]

EXAMPLE_ROW = (
    "示例（仅示范格式，别抄数值！）：\n"
    "NaCl-H2O | 305.15 | 0.313700 | 1.31 | EXP:NACL:EXAMPLE_2020 | "
    "\"Example Author, J. Example Chem., 1, 1 (2020), DOI: 10.0000/example\" | "
    "experimental | example2020::table1-row1 | 2026-08-16 | \"Table 1, p.2; 原文 10^-5 cm^2/s 直接作为 x\"\n\n"
    "⚠️ 只填真实核对过的数据；看不清/算不准的数值就空着并在 notes 写 UNVERIFIED，绝不编造。"
)

RULES = (
    "填写规则速记：\n"
    "1. 每个体系一个文件：把本模板另存为 candidates\\CaCl2候选.xlsx 等（工作表名保持 Confirmed_Records）；\n"
    "2. 一条实验记录一行；同一论文同一条件可多行（脚本自动取中位数）；\n"
    "3. 只录互扩散（mutual/interdiffusion），不要自扩散(self-)、示踪、模拟、拟合值；\n"
    "4. 排除 45 谱系清单里的文献（GPT检索提示词.md 附录 A / 检索方案第 2.2 节）；\n"
    "5. 填完把文件放进 py\\微调\\external\\candidates\\，跑 precheck 自查；\n"
    "6. 拿不准的体系/条件先问助手，不要硬录。"
)


def main() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Confirmed_Records"
    sheet.append(HEADERS)
    for index, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.font = Font(bold=True)
        sheet.column_dimensions[get_column_letter(index)].width = max(18, len(header) + 8)

    guide = workbook.create_sheet("填写说明")
    for row in EXPLANATIONS:
        guide.append(list(row))
    guide.cell(row=1, column=1).font = Font(bold=True)
    guide.column_dimensions["A"].width = 22
    guide.column_dimensions["B"].width = 8
    guide.column_dimensions["C"].width = 110
    guide.append([])
    guide.append([EXAMPLE_ROW])
    guide.append([])
    guide.append([RULES])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    if OUT.exists():
        raise FileExistsError(f"{OUT} exists; refusing to overwrite the template")
    workbook.save(OUT)
    print("template written:", OUT)
    print("sheets:", workbook.sheetnames)


if __name__ == "__main__":
    main()
