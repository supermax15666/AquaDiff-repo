from __future__ import annotations

import csv
import io
import math
import os
import statistics
import tempfile
from collections import defaultdict
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sklearn.model_selection import GroupKFold

from data_builder import SYSTEM_ORDER
from ml_common import (
    atomic_write_bytes,
    atomic_write_json,
    iter_jsonl,
    read_json,
    sha256_file,
    stable_hash,
    utc_now,
)
from models import MODEL_SPECS, NUMERIC_FEATURES, build_model


def _percentile(sorted_values: list[float], probability: float) -> float | None:
    if not sorted_values:
        return None
    if len(sorted_values) == 1:
        return sorted_values[0]
    position = (len(sorted_values) - 1) * probability
    lower = math.floor(position)
    upper = math.ceil(position)
    if lower == upper:
        return sorted_values[lower]
    weight = position - lower
    return sorted_values[lower] * (1 - weight) + sorted_values[upper] * weight


def _metrics_from(records: list[tuple[float, float, str]]) -> dict[str, Any]:
    if not records:
        raise ValueError("metric aggregation requires at least one record")
    apes: list[float] = []
    errors: list[float] = []
    squared: list[float] = []
    biases: list[float] = []
    lineage_apes: dict[str, list[float]] = defaultdict(list)
    for predicted, observed, lineage in records:
        error = predicted - observed
        ape = abs(error) / abs(observed) * 100.0
        errors.append(abs(error))
        squared.append(error * error)
        biases.append(error)
        apes.append(ape)
        lineage_apes[str(lineage)].append(ape)
    lineage_mdapes = [statistics.median(values) for values in lineage_apes.values()]
    n = len(records)
    return {
        "n": n,
        "coverage": 1.0,
        "mae_1e9_m2_s": statistics.fmean(errors),
        "rmse_1e9_m2_s": math.sqrt(statistics.fmean(squared)),
        "bias_1e9_m2_s": statistics.fmean(biases),
        "mdape_percent": statistics.median(apes),
        "macro_lineage_mdape_percent": statistics.fmean(lineage_mdapes),
        "n_lineages": len(lineage_apes),
        "accuracy_5pct": statistics.fmean(1.0 if ape <= 5.0 else 0.0 for ape in apes),
        "accuracy_10pct": statistics.fmean(1.0 if ape <= 10.0 else 0.0 for ape in apes),
        "p95_ape_percent": _percentile(sorted(apes), 0.95),
        "max_ape_percent": max(apes),
    }


def _fold_metrics(records: list[tuple[float, float, str]]) -> dict[str, Any]:
    full = _metrics_from(records)
    return {key: full[key] for key in ("n", "mae_1e9_m2_s", "mdape_percent", "accuracy_10pct")}


def _system_matrix(rows: list[dict[str, Any]], system: str) -> tuple[np.ndarray, np.ndarray, list[str], list[tuple[float, float, str]]]:
    selected = [row for row in rows if row["system"] == system]
    matrix = np.array(
        [[float(row[feature]) for feature in NUMERIC_FEATURES] for row in selected]
    )
    targets = np.array([float(row["diffusion_1e9_m2_s"]) for row in selected])
    groups = [str(row["origin_lineage_id"]) for row in selected]
    records = [
        (
            float(row["diffusion_1e9_m2_s"]),
            float(row["diffusion_1e9_m2_s"]),
            str(row["origin_lineage_id"]),
        )
        for row in selected
    ]
    return matrix, targets, groups, records


def protocol_fingerprint(config: dict[str, Any], dataset_sha256: str, code_dir: Path) -> str:
    material = {
        "schema_version": "ml_backend_d/v1",
        "config": {
            key: value
            for key, value in config.items()
            if key != "protocol_frozen"
        },
        "dataset_sha256": dataset_sha256,
        "models_sha256": sha256_file(code_dir / "models.py"),
        "evaluate_sha256": sha256_file(code_dir / "evaluate.py"),
        "data_builder_sha256": sha256_file(code_dir / "data_builder.py"),
    }
    return stable_hash(material)


def run_evaluation(
    config_path: Path,
    config: dict[str, Any],
    dataset_path: Path,
    output_root: Path,
) -> dict[str, Any]:
    rows = list(iter_jsonl(dataset_path))
    by_system: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        by_system[str(row["system"])].append(row)

    n_folds = int(config["cv"]["n_folds"])
    model_keys = [str(key) for key in config["model_keys"]]
    code_dir = config_path.parent
    fingerprint = protocol_fingerprint(config, sha256_file(dataset_path), code_dir)
    evaluation_dir = output_root / str(config["evaluation_dir"]) / fingerprint
    evaluation_dir.mkdir(parents=True, exist_ok=False)

    specialized: dict[str, dict[str, dict[str, Any]]] = {}
    pooled: dict[str, dict[str, dict[str, Any]]] = {}
    fold_details: dict[str, dict[str, list[dict[str, Any]]]] = {}

    system_folds: dict[str, list[tuple[np.ndarray, np.ndarray]]] = {}
    system_matrices: dict[str, dict[str, Any]] = {}
    for system in SYSTEM_ORDER:
        selected = by_system.get(system, [])
        if not selected:
            continue
        matrix = np.array(
            [[float(row[feature]) for feature in NUMERIC_FEATURES] for row in selected]
        )
        targets = np.array([float(row["diffusion_1e9_m2_s"]) for row in selected])
        groups = [str(row["origin_lineage_id"]) for row in selected]
        splitter = GroupKFold(n_splits=n_folds)
        folds = [(train, test) for train, test in splitter.split(matrix, targets, groups)]
        system_folds[system] = folds
        system_matrices[system] = {
            "matrix": matrix,
            "targets": targets,
            "groups": groups,
        }
        fold_details[system] = {model_key: [] for model_key in model_keys}

    for system in SYSTEM_ORDER:
        if system not in system_matrices:
            continue
        matrix = system_matrices[system]["matrix"]
        targets = system_matrices[system]["targets"]
        groups = system_matrices[system]["groups"]
        specialized[system] = {}
        for model_key in model_keys:
            collected: list[tuple[float, float, str]] = []
            for fold_index, (train, test) in enumerate(system_folds[system]):
                model = build_model(model_key, pooled=False)
                model.fit(matrix[train], targets[train])
                predicted = model.predict(matrix[test])
                for index in range(len(test)):
                    collected.append(
                        (float(predicted[index]), float(targets[test[index]]), groups[test[index]])
                    )
                fold_records = [
                    (
                        float(predicted[index]),
                        float(targets[test[index]]),
                        groups[test[index]],
                    )
                    for index in range(len(test))
                ]
                fold_details[system][model_key].append(_fold_metrics(fold_records))
            specialized[system][model_key] = _metrics_from(collected)

    if bool(config.get("pooled_control", True)):
        all_features: list[list[Any]] = []
        all_targets: list[float] = []
        all_system_labels: list[str] = []
        for system in SYSTEM_ORDER:
            if system not in system_matrices:
                continue
            selected = by_system.get(system, [])
            for row in selected:
                all_features.append([float(row[f]) for f in NUMERIC_FEATURES])
                all_system_labels.append(system)
                all_targets.append(float(row["diffusion_1e9_m2_s"]))
        feature_matrix = np.array(all_features)
        target_array = np.array(all_targets)
        system_array = np.array(all_system_labels)
        pooled_matrix = np.hstack(
            [system_array.reshape(-1, 1).astype(object), feature_matrix]
        )
        position = {system: index for index, system in enumerate(SYSTEM_ORDER)}
        row_system_indices: dict[str, list[int]] = defaultdict(list)
        for global_index, system in enumerate(all_system_labels):
            row_system_indices[system].append(global_index)
        for system in SYSTEM_ORDER:
            if system not in system_matrices:
                continue
            local_indices = row_system_indices[system]
            pooled[system] = {}
            for model_key in model_keys:
                collected: list[tuple[float, float, str]] = []
                for fold_index, (train, test) in enumerate(system_folds[system]):
                    train_global: list[int] = []
                    for other_system in SYSTEM_ORDER:
                        if other_system not in system_folds:
                            continue
                        other_train, _other_test = system_folds[other_system][fold_index]
                        train_global.extend(
                            row_system_indices[other_system][int(flat)]
                            for flat in other_train
                        )
                    model = build_model(model_key, pooled=True)
                    model.fit(pooled_matrix[train_global], target_array[train_global])
                    test_global = [
                        row_system_indices[system][int(flat)] for flat in test
                    ]
                    predicted = model.predict(pooled_matrix[test_global])
                    local_groups = system_matrices[system]["groups"]
                    for offset, global_index in enumerate(test_global):
                        collected.append(
                            (
                                float(predicted[offset]),
                                target_array[global_index],
                                local_groups[int(test[offset])],
                            )
                        )
                pooled[system][model_key] = _metrics_from(collected)

    primary_metric = str(config["selection"]["primary_metric"])
    tie_break = str(config["selection"]["tie_break_metric"])
    selection: dict[str, Any] = {}
    for system in SYSTEM_ORDER:
        if system not in specialized:
            continue
        ranked = sorted(
            specialized[system],
            key=lambda key: (
                specialized[system][key][primary_metric],
                specialized[system][key][tie_break],
                key,
            ),
        )
        selected_key = ranked[0]
        selection[system] = {
            "selected_model_key": selected_key,
            "selected_model_display": MODEL_SPECS[selected_key]["display_name"],
            "primary_metric": primary_metric,
            "primary_metric_value": specialized[system][selected_key][primary_metric],
            "tie_break_metric": tie_break,
            "tie_break_value": specialized[system][selected_key][tie_break],
            "ranking": ranked,
        }

    pooled_delta: dict[str, dict[str, float | None]] = {}
    if pooled:
        for system in SYSTEM_ORDER:
            if system not in selection or system not in pooled:
                continue
            key = selection[system]["selected_model_key"]
            pooled_delta[system] = {
                "specialized": specialized[system][key][primary_metric],
                "pooled": pooled[system][key][primary_metric],
                "delta_specialized_minus_pooled": (
                    specialized[system][key][primary_metric]
                    - pooled[system][key][primary_metric]
                ),
            }

    results: dict[str, Any] = {
        "schema_version": "ml_backend_d/v1",
        "artifact_type": "ml_backend_d_evaluation",
        "created_at_utc": utc_now(),
        "protocol_fingerprint": fingerprint,
        "protocol_frozen": bool(config.get("protocol_frozen", False)),
        "dataset_path": str(dataset_path),
        "dataset_sha256": sha256_file(dataset_path),
        "cv": dict(config["cv"]),
        "model_specs": {
            key: {
                "display_name": MODEL_SPECS[key]["display_name"],
                "description": MODEL_SPECS[key]["description"],
            }
            for key in model_keys
        },
        "primary_metric": primary_metric,
        "specialized": specialized,
        "pooled": pooled,
        "selection": selection,
        "pooled_delta": pooled_delta,
        "fold_details": fold_details,
        "code": {
            "models_sha256": sha256_file(code_dir / "models.py"),
            "evaluate_sha256": sha256_file(code_dir / "evaluate.py"),
            "data_builder_sha256": sha256_file(code_dir / "data_builder.py"),
        },
    }
    atomic_write_json(evaluation_dir / "results.json", results)

    csv_rows: list[dict[str, Any]] = []
    for system in SYSTEM_ORDER:
        for model_key in model_keys:
            if system in specialized and model_key in specialized[system]:
                csv_rows.append(
                    _csv_row(system, model_key, "specialized", specialized[system][model_key])
                )
            if system in pooled and model_key in pooled[system]:
                csv_rows.append(_csv_row(system, model_key, "pooled", pooled[system][model_key]))
    fieldnames = [
        "system",
        "model_key",
        "model_display",
        "role",
        "n",
        "n_lineages",
        "coverage",
        "mae_1e9_m2_s",
        "rmse_1e9_m2_s",
        "bias_1e9_m2_s",
        "mdape_percent",
        "macro_lineage_mdape_percent",
        "accuracy_5pct",
        "accuracy_10pct",
        "p95_ape_percent",
        "max_ape_percent",
    ]
    csv_path = evaluation_dir / "manuscript_metrics.csv"
    descriptor, temp_name = tempfile.mkstemp(
        prefix=".manuscript.", suffix=".tmp", dir=str(evaluation_dir)
    )
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(csv_rows)
        os.replace(temp_name, csv_path)
    except BaseException:
        Path(temp_name).unlink(missing_ok=True)
        raise

    write_report(evaluation_dir, results)
    return results


def _csv_row(
    system: str, model_key: str, role: str, metrics: dict[str, Any]
) -> dict[str, Any]:
    return {
        "system": system,
        "model_key": model_key,
        "model_display": MODEL_SPECS[model_key]["display_name"],
        "role": role,
        **metrics,
    }


def write_report(evaluation_dir: Path, results: dict[str, Any]) -> None:
    lines = ["# ML Backend D: 四体系数值模型比较", ""]
    lines.append("谱系分组 5 折交叉验证；选模主指标为谱系宏平均 MdAPE。")
    lines.append("")
    lines.append(
        "| 体系 | 入选模型 | 宏谱系 MdAPE% | MAE | Acc@10% | pooled 对照 MdAPE% | Δ(spec−pooled) |"
    )
    lines.append("|---|---|---|---:|---:|---:|---:|")
    for system in SYSTEM_ORDER:
        selection = results["selection"].get(system)
        if not selection:
            continue
        key = selection["selected_model_key"]
        spec = results["specialized"][system][key]
        delta = results["pooled_delta"].get(system, {})
        lines.append(
            f"| {system} | {selection['selected_model_display']} "
            f"({key}) | {spec['macro_lineage_mdape_percent']:.2f} | "
            f"{spec['mae_1e9_m2_s']:.4f} | {spec['accuracy_10pct'] * 100:.1f}% | "
            f"{delta.get('pooled', 0.0):.2f} | {delta.get('delta_specialized_minus_pooled', 0.0):+.2f} |"
        )
    lines.append("")
    lines.append("全部模型明细见 manuscript_metrics.csv 与 results.json。")
    lines.append("")
    report_path = evaluation_dir / "report.md"
    report_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


# ---------------------------------------------------------------------------
# 人类可读汇总导出：冻结汇总.csv（权威，UTF-8 BOM）与 冻结汇总.xlsx（阅读用）
# ---------------------------------------------------------------------------

SUMMARY_CSV_NAME = "冻结汇总.csv"
SUMMARY_XLSX_NAME = "冻结汇总.xlsx"

_SUMMARY_METRIC_KEYS = [
    "n",
    "n_lineages",
    "coverage",
    "mae_1e9_m2_s",
    "rmse_1e9_m2_s",
    "bias_1e9_m2_s",
    "mdape_percent",
    "macro_lineage_mdape_percent",
    "accuracy_5pct",
    "accuracy_10pct",
    "p95_ape_percent",
    "max_ape_percent",
]

_SUMMARY_CSV_FIELDNAMES = [
    "system",
    "model_key",
    "model_display",
    "role",
    "selected",
    *_SUMMARY_METRIC_KEYS,
    "protocol_fingerprint",
    "protocol_frozen",
    "dataset_sha256",
]


def _summary_rows(results: dict[str, Any]) -> list[dict[str, Any]]:
    model_keys = list(results.get("model_specs", {}).keys())
    rows: list[dict[str, Any]] = []
    for system in SYSTEM_ORDER:
        selected_key = (
            results.get("selection", {}).get(system, {}).get("selected_model_key")
        )
        for model_key in model_keys:
            for role in ("specialized", "pooled"):
                metrics = results.get(role, {}).get(system, {}).get(model_key)
                if not metrics:
                    continue
                rows.append(
                    {
                        "system": system,
                        "model_key": model_key,
                        "model_display": results["model_specs"][model_key][
                            "display_name"
                        ],
                        "role": role,
                        "selected": (
                            "是"
                            if role == "specialized" and model_key == selected_key
                            else ""
                        ),
                        **{key: metrics[key] for key in _SUMMARY_METRIC_KEYS},
                        "protocol_fingerprint": str(results["protocol_fingerprint"]),
                        "protocol_frozen": bool(results.get("protocol_frozen", False)),
                        "dataset_sha256": str(results["dataset_sha256"]),
                    }
                )
    return rows


def _write_summary_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    buffer = io.StringIO()
    writer = csv.DictWriter(
        buffer,
        fieldnames=_SUMMARY_CSV_FIELDNAMES,
        extrasaction="ignore",
        lineterminator="\n",
    )
    writer.writeheader()
    writer.writerows(rows)
    atomic_write_bytes(path, buffer.getvalue().encode("utf-8-sig"))


def _excel_style_header(cell: Any) -> None:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="4472C4")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _excel_style_selected(cell: Any) -> None:
    cell.fill = PatternFill("solid", fgColor="C6EFCE")
    cell.font = Font(bold=True)


def _write_table(
    sheet: Any,
    headers: list[str],
    rows: list[list[Any]],
    *,
    formats: dict[int, str] | None = None,
    widths: list[float] | None = None,
) -> None:
    for column_index, header in enumerate(headers, start=1):
        _excel_style_header(sheet.cell(row=1, column=column_index, value=header))
        if widths is not None and column_index <= len(widths):
            sheet.column_dimensions[get_column_letter(column_index)].width = widths[
                column_index - 1
            ]
    for row_index, row in enumerate(rows, start=2):
        for column_index, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            if formats and column_index in formats:
                cell.number_format = formats[column_index]


def _write_notes_sheet(wb: Workbook, results: dict[str, Any], registry: dict[str, Any],
                       evaluation_dir: Path, registry_dir: Path) -> None:
    notes = [
        (
            "文件角色",
            "本 Excel 仅供阅读；数值一律以《冻结汇总.csv》为准；逐位权威数据见 "
            f"{evaluation_dir / 'results.json'} 与 {registry_dir / 'registry.json'}。",
        ),
        ("协议指纹", str(results["protocol_fingerprint"])),
        (
            "协议冻结状态",
            "已冻结（protocol_frozen=true，论文主表可用）"
            if results.get("protocol_frozen")
            else "未冻结（protocol_frozen=false，仅供探索性参考）",
        ),
        ("生成时间（UTC）", str(results.get("created_at_utc", ""))),
        ("数据集 SHA-256", str(results["dataset_sha256"])),
        ("评价目录", str(evaluation_dir)),
        ("注册目录", str(registry_dir)),
        (
            "数据",
            "642 个合格唯一条件：NaCl 174 / KCl 247 / LiCl 132 / CaCl2 89；"
            "同一 (system, T, c) 的多条文献记录取中位数共识值；与 642 题 evaluator gold 精确一致。",
        ),
        (
            "训练/测试划分",
            "每体系独立做 5 折 GroupKFold，按 origin_lineage_id 整谱系分组、不 shuffle："
            "同一文献谱系的全部条件要么整组在训练折、要么整组在测试折，防止近重复条件泄漏。"
            "每折约 4/5 谱系训练、1/5 谱系测试；五个折的测试预测聚合后统一计算指标"
            "（每个条件恰好得到一次留出预测），六个候选模型使用完全相同的折划分。",
        ),
        (
            "候选模型",
            "Polynomial Ridge / RBF-SVR / Random Forest / Gradient Boosting / "
            "Gaussian Process / Dummy（训练中位数基线），超参与随机种子预先冻结，评估可逐位复现。",
        ),
        (
            "选模规则",
            "主指标 = 谱系宏平均 MdAPE%（每个谱系内取 APE 中位数，再跨谱系等权平均），"
            "tie-break = MAE；每体系注册 1 个专用模型，并在该体系全部数据上重训后交付。",
        ),
        (
            "pooled 对照",
            "复用同一折划分：训练 = 本体系训练折 + 其余三体系的同折训练折（system 作 one-hot 特征），"
            "仅在本体系测试折上预测，与专用模型同题配对；只用于比较，不作为 Agent 后端。",
        ),
        (
            "适用域与拒绝",
            "温度×浓度矩形盒（各体系 [min, max]）；盒外输入一律结构化拒绝，"
            "不得用其他体系模型代替。",
        ),
        (
            "证据等级",
            "回顾性、探索性研究；不得表述为盲测、前瞻性外部验证或确认性结果。",
        ),
        (
            "LiCl 提示",
            "132 条件中 108 条来自单一谱系 REF_087；该谱系作测试折时训练侧仅 24 条件，"
            "入选 GB 的整体 MdAPE 为 19.5%（宏谱系均值 4.68%），正文需同时报告两者。",
        ),
        (
            "折级明细",
            "『折级明细』表仅含专用模型五折明细（pooled 对照按同折配对，不单独记录折级）。"
            "论文图 10 的置信区间需另行补充谱系聚类 bootstrap。",
        ),
    ]
    sheet = wb.active
    sheet.title = "说明"
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 115
    for row_index, (label, value) in enumerate(notes, start=1):
        label_cell = sheet.cell(row=row_index, column=1, value=label)
        label_cell.font = Font(bold=True)
        label_cell.alignment = Alignment(vertical="top")
        value_cell = sheet.cell(row=row_index, column=2, value=value)
        value_cell.alignment = Alignment(vertical="top", wrap_text=True)


def _write_selected_sheet(wb: Workbook, results: dict[str, Any],
                          registry: dict[str, Any]) -> None:
    headers = [
        "体系", "入选模型", "宏谱系MdAPE%(专用)", "MAE", "RMSE", "bias",
        "整体MdAPE%", "Acc@5%", "Acc@10%", "P95APE%", "maxAPE%",
        "训练条件数", "谱系数", "T下限K", "T上限K", "c下限mol/L", "c上限mol/L",
        "pooled宏MdAPE%", "Δ(专用−pooled)",
    ]
    formats = {
        3: "0.00", 4: "0.0000", 5: "0.0000", 6: "0.0000", 7: "0.00",
        8: "0.0%", 9: "0.0%", 10: "0.00", 11: "0.00",
        14: "0.0", 15: "0.0", 16: "0.00000", 17: "0.00000",
        18: "0.00", 19: "+0.00;-0.00;0.00",
    }
    rows: list[list[Any]] = []
    for entry in registry.get("systems", []):
        system = str(entry["system"])
        key = str(entry["registered_model_key"])
        spec = results["specialized"][system][key]
        pooled = results["pooled"][system][key]
        domain = entry["applicability_domain"]
        rows.append(
            [
                system,
                entry["registered_model_display"],
                spec["macro_lineage_mdape_percent"],
                spec["mae_1e9_m2_s"],
                spec["rmse_1e9_m2_s"],
                spec["bias_1e9_m2_s"],
                spec["mdape_percent"],
                spec["accuracy_5pct"],
                spec["accuracy_10pct"],
                spec["p95_ape_percent"],
                spec["max_ape_percent"],
                entry["trained_on_conditions"],
                entry["trained_on_lineages"],
                domain["temperature_K"]["min"],
                domain["temperature_K"]["max"],
                domain["concentration_mol_L"]["min"],
                domain["concentration_mol_L"]["max"],
                pooled["macro_lineage_mdape_percent"],
                spec["macro_lineage_mdape_percent"]
                - pooled["macro_lineage_mdape_percent"],
            ]
        )
    sheet = wb.create_sheet("入选模型")
    _write_table(sheet, headers, rows, formats=formats)
    sheet.freeze_panes = "C2"


def _write_grid_sheet(
    wb: Workbook,
    title: str,
    results: dict[str, Any],
    role: str,
    systems: list[str],
    model_keys: list[str],
) -> None:
    sheet = wb.create_sheet(title)
    displays = [results["model_specs"][key]["display_name"] for key in model_keys]
    selection = results.get("selection", {})

    def add_grid(
        start_row: int,
        caption: str,
        get_value: Any,
        number_format: str,
        highlight_selected: bool,
    ) -> int:
        sheet.cell(row=start_row, column=1, value=caption).font = Font(bold=True)
        header_row = start_row + 1
        _excel_style_header(sheet.cell(row=header_row, column=1, value="体系"))
        for column_index, display in enumerate(displays, start=2):
            _excel_style_header(
                sheet.cell(row=header_row, column=column_index, value=display)
            )
        for row_offset, system in enumerate(systems):
            row_index = header_row + 1 + row_offset
            system_cell = sheet.cell(row=row_index, column=1, value=system)
            system_cell.font = Font(bold=True)
            for column_offset, model_key in enumerate(model_keys, start=2):
                cell = sheet.cell(
                    row=row_index,
                    column=column_offset,
                    value=get_value(system, model_key),
                )
                cell.number_format = number_format
                if (
                    highlight_selected
                    and selection.get(system, {}).get("selected_model_key") == model_key
                ):
                    _excel_style_selected(cell)
        return header_row + 1 + len(systems)

    cursor = 1
    cursor = add_grid(
        cursor,
        "宏谱系 MdAPE%（主选模指标；绿底 = 入选注册模型）",
        lambda system, key: results[role][system][key]["macro_lineage_mdape_percent"],
        "0.00",
        highlight_selected=(role == "specialized"),
    )
    cursor += 1
    cursor = add_grid(
        cursor,
        "CV Accuracy@10%（全部测试预测）",
        lambda system, key: results[role][system][key]["accuracy_10pct"],
        "0.0%",
        highlight_selected=(role == "specialized"),
    )
    cursor += 1
    cursor = add_grid(
        cursor,
        "MAE（10^-9 m^2/s）",
        lambda system, key: results[role][system][key]["mae_1e9_m2_s"],
        "0.0000",
        highlight_selected=(role == "specialized"),
    )
    sheet.column_dimensions["A"].width = 26
    for column_index in range(2, 2 + len(displays)):
        sheet.column_dimensions[get_column_letter(column_index)].width = 22


def _write_paired_sheet(wb: Workbook, results: dict[str, Any],
                        registry: dict[str, Any]) -> None:
    headers = ["体系", "入选模型家族", "专用宏MdAPE%", "pooled宏MdAPE%", "Δ(专用−pooled)", "解读"]
    formats = {3: "0.00", 4: "0.00", 5: "+0.00;-0.00;0.00"}
    rows: list[list[Any]] = []
    for entry in registry.get("systems", []):
        system = str(entry["system"])
        key = str(entry["registered_model_key"])
        spec = results["specialized"][system][key]["macro_lineage_mdape_percent"]
        pooled = results["pooled"][system][key]["macro_lineage_mdape_percent"]
        delta = spec - pooled
        rows.append(
            [
                system,
                results["model_specs"][key]["display_name"],
                spec,
                pooled,
                delta,
                "pooled 更优（数据借用有利）" if delta > 0 else "专用模型更优",
            ]
        )
    sheet = wb.create_sheet("专用vs pooled")
    _write_table(sheet, headers, rows, formats=formats)
    sheet.freeze_panes = "A2"
    note = sheet.cell(
        row=2 + len(rows),
        column=1,
        value="注：同折配对（pooled 训练 = 本体系训练折 + 其余三体系同折训练折），"
        "Δ<0 表示专用模型宏谱系 MdAPE 更低。",
    )
    note.font = Font(italic=True)


def _write_detail_sheet(wb: Workbook, rows: list[dict[str, Any]]) -> None:
    headers = [
        "体系", "模型", "角色", "入选", "n", "谱系数", "覆盖率",
        "MAE", "RMSE", "bias", "整体MdAPE%", "宏谱系MdAPE%",
        "Acc@5%", "Acc@10%", "P95APE%", "maxAPE%",
    ]
    formats = {
        7: "0.0%", 8: "0.0000", 9: "0.0000", 10: "0.0000",
        11: "0.00", 12: "0.00", 13: "0.0%", 14: "0.0%",
        15: "0.00", 16: "0.00",
    }
    table_rows: list[list[Any]] = []
    for row in rows:
        table_rows.append(
            [
                row["system"],
                row["model_display"],
                "专用" if row["role"] == "specialized" else "pooled",
                row["selected"],
                row["n"],
                row["n_lineages"],
                row["coverage"],
                row["mae_1e9_m2_s"],
                row["rmse_1e9_m2_s"],
                row["bias_1e9_m2_s"],
                row["mdape_percent"],
                row["macro_lineage_mdape_percent"],
                row["accuracy_5pct"],
                row["accuracy_10pct"],
                row["p95_ape_percent"],
                row["max_ape_percent"],
            ]
        )
    sheet = wb.create_sheet("全部明细")
    _write_table(sheet, headers, table_rows, formats=formats)
    sheet.freeze_panes = "A2"


def _write_fold_sheet(wb: Workbook, results: dict[str, Any]) -> None:
    headers = ["体系", "模型", "折", "n", "MAE", "MdAPE%", "Acc@10%"]
    formats = {5: "0.0000", 6: "0.00", 7: "0.0%"}
    rows: list[list[Any]] = []
    model_keys = list(results.get("model_specs", {}).keys())
    displays = [results["model_specs"][key]["display_name"] for key in model_keys]
    for system in SYSTEM_ORDER:
        for key_index, model_key in enumerate(model_keys):
            folds = results.get("fold_details", {}).get(system, {}).get(model_key, [])
            for fold_index, fold in enumerate(folds, start=1):
                rows.append(
                    [
                        system,
                        displays[key_index],
                        fold_index,
                        fold["n"],
                        fold["mae_1e9_m2_s"],
                        fold["mdape_percent"],
                        fold["accuracy_10pct"],
                    ]
                )
    sheet = wb.create_sheet("折级明细")
    _write_table(sheet, headers, rows, formats=formats)
    sheet.freeze_panes = "A2"


def export_human_summary(
    output_root: Path,
    evaluation_dir: Path,
    registry_dir: Path,
) -> dict[str, str]:
    """由 register 调用：把当前评价/注册结果导出为人类可读 CSV 与 Excel。

    CSV（冻结汇总.csv）为权威汇总（UTF-8 BOM，Excel 可直接打开）；
    Excel（冻结汇总.xlsx）仅作阅读用。逐位权威数据仍以 evaluation/registry
    指纹目录下的 results.json / registry.json 为准。
    """
    results = read_json(evaluation_dir / "results.json")
    registry = read_json(registry_dir / "registry.json")
    rows = _summary_rows(results)

    csv_path = output_root / SUMMARY_CSV_NAME
    _write_summary_csv(csv_path, rows)

    wb = Workbook()
    _write_notes_sheet(wb, results, registry, evaluation_dir, registry_dir)
    _write_selected_sheet(wb, results, registry)
    model_keys = list(results.get("model_specs", {}).keys())
    _write_grid_sheet(
        wb, "专用模型对比", results, "specialized", SYSTEM_ORDER, model_keys
    )
    _write_grid_sheet(
        wb, "pooled对照", results, "pooled", SYSTEM_ORDER, model_keys
    )
    _write_paired_sheet(wb, results, registry)
    _write_detail_sheet(wb, rows)
    _write_fold_sheet(wb, results)

    xlsx_path = output_root / SUMMARY_XLSX_NAME
    descriptor, temp_name = tempfile.mkstemp(
        prefix=".冻结汇总.", suffix=".tmp", dir=str(output_root)
    )
    try:
        os.close(descriptor)
        wb.save(temp_name)
        os.replace(temp_name, xlsx_path)
    except BaseException:
        Path(temp_name).unlink(missing_ok=True)
        raise
    return {"csv": str(csv_path), "xlsx": str(xlsx_path)}
